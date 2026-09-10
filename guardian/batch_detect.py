#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""批量文件检测。

给定一组文件 / 目录，逐个抽取文本跑新内核检测，汇总风险分布并导出 CSV。
设计要点：

* 纯 stdlib（csv）+ 内核层（engine / schema），不依赖 tkinter / fastapi。
* 文本抽取：.txt/.md/.csv/.json/.html/.py/.log 等按文本读取（UTF-8，Windows 回退 GBK）；
  .docx 在 python-docx 可用时抽取正文，否则跳过并标记。
* 聚合：每文件风险等级 / 分数 / 各级命中数；整体汇总（文件数、违规文件数、
  严重度分布、Top 风险词）。
* 导出 CSV（stdlib）便于在 Excel 打开复核；Excel(xlsx) 在 openpyxl 可用时一并生成。
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from guardian.engine import DetectionEngine
from guardian.schema import DetectionOptions

_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".json", ".html",
                  ".htm", ".xml", ".py", ".log", ".text", ".rst"}
_DOCX_SUFFIXES = {".docx"}


@dataclass
class BatchFileResult:
    path: str
    text_len: int = 0
    risk_level: str = "基本合规"
    score: int = 100
    counts: dict = field(default_factory=dict)
    findings_count: int = 0
    top_findings: list = field(default_factory=list)
    skipped: bool = False
    error: str = ""


@dataclass
class BatchSummary:
    total_files: int = 0
    scanned: int = 0
    skipped: int = 0
    risky_files: int = 0
    severity_dist: dict = field(default_factory=dict)
    top_keywords: dict = field(default_factory=dict)


def _read_text_file(path: Path) -> str:
    """读取文本文件，UTF-8 优先，Windows 常见 GBK 回退。"""
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _extract_text(path: Path) -> str:
    """从文件抽取纯文本；不支持的类型返回空字符串（由调用方决定跳过）。"""
    suffix = path.suffix.lower()
    if suffix in _TEXT_SUFFIXES:
        return _read_text_file(path)
    if suffix in _DOCX_SUFFIXES:
        try:
            from docx import Document
        except ImportError:
            return ""
        try:
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text)
        except Exception:
            return ""
    return ""


class BatchRunner:
    """批量检测编排器。"""

    def __init__(self, engine: Optional[DetectionEngine] = None,
                 options: Optional[DetectionOptions] = None):
        self.engine = engine or DetectionEngine()
        self.options = options or DetectionOptions()

    # -------------------------------------------------- 扫描

    def scan(self, paths: list, recursive: bool = False) -> list[BatchFileResult]:
        files: list[Path] = []
        for p in paths:
            pp = Path(p)
            if pp.is_dir():
                if recursive:
                    files.extend(sorted(pp.rglob("*")))
                else:
                    files.extend(sorted(pp.iterdir()))
            elif pp.is_file():
                files.append(pp)

        results: list[BatchFileResult] = []
        supported = _TEXT_SUFFIXES | _DOCX_SUFFIXES
        for f in files:
            if f.is_dir():
                continue
            if f.suffix.lower() not in supported:
                # 不支持的类型：作为"跳过"项进入报告，便于用户在汇总里看到
                results.append(BatchFileResult(
                    path=str(f), skipped=True, error="不支持的文件类型"))
                continue
            results.append(self._scan_one(f))
        return results

    def _scan_one(self, path: Path) -> BatchFileResult:
        text = _extract_text(path)
        if not text.strip():
            return BatchFileResult(path=str(path), skipped=True,
                                   error="空文件或无法抽取文本")
        try:
            r = self.engine.detect_text(text, self.options)
        except Exception as e:  # 单文件失败不影响整体
            return BatchFileResult(path=str(path), skipped=True, error=str(e)[:120])
        return BatchFileResult(
            path=str(path),
            text_len=len(text),
            risk_level=r.summary["risk_level"],
            score=r.summary["score"],
            counts=r.summary["counts"],
            findings_count=len(r.findings),
            top_findings=[f.matched_text or f.keyword
                          for f in r.findings[:5]],
        )

    # -------------------------------------------------- 汇总

    def summarize(self, results: list[BatchFileResult]) -> BatchSummary:
        s = BatchSummary()
        s.total_files = len(results)
        s.scanned = sum(1 for r in results if not r.skipped)
        s.skipped = sum(1 for r in results if r.skipped)
        s.risky_files = sum(1 for r in results
                            if not r.skipped and r.risk_level != "基本合规")
        sev = {"critical": 0, "high": 0, "medium": 0, "low": 0}
        kw: dict[str, int] = {}
        for r in results:
            if r.skipped:
                continue
            for k, v in (r.counts or {}).items():
                sev[k] = sev.get(k, 0) + v
            for k in r.top_findings:
                kw[k] = kw.get(k, 0) + 1
        s.severity_dist = sev
        s.top_keywords = dict(sorted(kw.items(), key=lambda x: -x[1])[:10])
        return s

    # -------------------------------------------------- 导出

    def export_csv(self, results: list[BatchFileResult],
                   out_path: str | Path) -> Path:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        cols = ["path", "risk_level", "score", "findings_count",
                "critical", "high", "medium", "low", "text_len",
                "top_findings", "skipped", "error"]
        with out_path.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(cols)
            for r in results:
                c = r.counts or {}
                w.writerow([
                    r.path, r.risk_level, r.score, r.findings_count,
                    c.get("critical", 0), c.get("high", 0),
                    c.get("medium", 0), c.get("low", 0), r.text_len,
                    " / ".join(r.top_findings), r.skipped, r.error,
                ])
        return out_path

    def export_xlsx(self, results: list[BatchFileResult],
                    out_path: str | Path) -> Optional[Path]:
        """在 openpyxl 可用时生成 xlsx；否则返回 None。"""
        try:
            from openpyxl import Workbook
        except ImportError:
            return None
        wb = Workbook()
        ws = wb.active
        ws.title = "批量检测"
        cols = ["路径", "风险等级", "合规分", "命中数", "高危", "中危",
                "低危", "提示", "字数", "Top命中", "跳过", "备注"]
        ws.append(cols)
        for r in results:
            c = r.counts or {}
            ws.append([
                r.path, r.risk_level, r.score, r.findings_count,
                c.get("critical", 0), c.get("high", 0), c.get("medium", 0),
                c.get("low", 0), r.text_len, " / ".join(r.top_findings),
                r.skipped, r.error,
            ])
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out_path))
        return out_path
